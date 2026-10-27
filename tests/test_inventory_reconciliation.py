import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from inventory_reconciliation import (MAX_FILE_BYTES, ReconciliationError,
    ReconciliationService, import_schema_statements, migrate_import_schema, parse_inventory_file)
from item_registry import Actor, Registry, RegistryError
from item_registry_schema import migrate_registry


WHEN = '2026-10-22T09:00:00+09:00'


def csv_bytes(value):
    return value.encode('utf-8-sig')


def fixture():
    conn = sqlite3.connect(':memory:')
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    conn.execute('CREATE TABLE users(id INTEGER PRIMARY KEY,role TEXT,admin_permissions TEXT)')
    conn.executemany('INSERT INTO users VALUES(?,?,?)', [(1, 'owner', None), (2, 'user', None),
        (3, 'user', None), (4, 'admin', '["documents"]')])
    conn.execute('''CREATE TABLE merchandise(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER REFERENCES users(id),
        product_name TEXT,purchase_price INTEGER DEFAULT 0,purchase_date TEXT,sale_date TEXT,sale_price INTEGER,
        is_listed INTEGER,is_shipped INTEGER,scope TEXT,custody_location TEXT)''')
    conn.executemany("INSERT INTO merchandise(id,user_id,product_name,purchase_price,scope) VALUES(?,?,?,?,'user')",
                     [(10, 2, '既存バッグ', 30000), (11, 3, '別顧客の靴', 20000)])
    migrate_registry(conn)
    migrate_import_schema(conn)
    conn.commit()
    return conn


class ParserTests(unittest.TestCase):
    def test_legacy_export_maps_known_columns_and_does_not_infer_owner_or_fees(self):
        p = parse_inventory_file(csv_bytes('ID,登録者,商品名,仕入額,利益,利益率,発送済み\n10,架空名,既存バッグ,30000,0,0%,\n'), 'data.csv')
        self.assertEqual(p['error_count'], 0)
        r = p['rows'][0]['record']
        self.assertEqual((r['legacy_id'], r['purchase_price']), (10, 30000))
        self.assertIsNone(r['owner_user_id'])
        self.assertIsNone(r['sale_price'])
        self.assertIsNone(r['is_shipped'])
        self.assertEqual(r['custody_location'], 'unknown')
        self.assertIn('登録者', p['unused_headers'])
        self.assertFalse(p['creates_charges'])
        self.assertFalse(p['creates_work_records'])

    def test_cp932_and_manual_mapping(self):
        p = parse_inventory_file('番号,名称,原価\na-1,架空の靴,"30,000"\n'.encode('cp932'), 'example.CSV',
                                 {'external_key': '番号', 'title': '名称', 'purchase_price': '原価'})
        self.assertEqual(p['encoding'], 'cp932')
        self.assertEqual(p['rows'][0]['record']['purchase_price'], 30000)

    def test_no_identifier_is_held_without_using_title_or_row_number(self):
        p = parse_inventory_file(csv_bytes('商品名\n架空バッグ\n'), 'file.csv')
        self.assertEqual(p['error_count'], 1)
        self.assertIn('行番号', ' '.join(p['rows'][0]['errors']))

    def test_duplicate_ids_and_external_keys_hold_every_ambiguous_row(self):
        p = parse_inventory_file(csv_bytes('ID,外部キー,商品名\n10,a,商品\n10,b,商品\n11,a,別商品\n'), 'file.csv')
        self.assertEqual(p['error_count'], 3)

    def test_missing_values_remain_unknown_not_zero(self):
        p = parse_inventory_file(csv_bytes('外部キー,商品名,仕入額,入庫日,登録完了日\na,仮の商品,,,\n'), 'f.csv')
        r = p['rows'][0]['record']
        self.assertIsNone(r['purchase_price'])
        self.assertIsNone(r['received_at'])
        self.assertIsNone(r['registration_completed_at'])

    def test_invalid_money_dates_quantity_and_boolean_are_row_errors(self):
        p = parse_inventory_file(csv_bytes('外部キー,商品名,仕入額,仕入日,数量,発送済み\na,商品,1.5,46000,2,たぶん\n'), 'f.csv')
        self.assertEqual(len(p['rows'][0]['errors']), 4)

    def test_strict_headers_mapping_and_file_limits(self):
        cases = [(b'', 'f.csv', None), (b'a' * (MAX_FILE_BYTES + 1), 'f.csv', None),
                 (csv_bytes('商品名,商品名\nA,B'), 'f.csv', None),
                 (csv_bytes('外部キー,商品名\na,A'), 'f.csv', {'title': '不明'}),
                 (b'bad', 'f.xlsm', None)]
        for data, filename, mapping in cases:
            with self.subTest(filename=filename, length=len(data)), self.assertRaises(ReconciliationError):
                parse_inventory_file(data, filename, mapping)

    def test_formula_csv_is_rejected_as_a_row_not_executed(self):
        p = parse_inventory_file(csv_bytes('外部キー,商品名\na,=WEBSERVICE("example")\n'), 'f.csv')
        self.assertEqual(p['error_count'], 1)

    def test_malformed_mapping_returns_validation_error_not_server_error(self):
        for mapping in ([], 'title', {'title': []}, {'title': None}):
            with self.subTest(mapping=mapping), self.assertRaises(ReconciliationError) as raised:
                parse_inventory_file(csv_bytes('外部キー,商品名\na,商品\n'), 'f.csv', mapping)
            self.assertEqual(raised.exception.code, 'mapping')

    def make_xlsx(self, rows):
        from openpyxl import Workbook
        book = Workbook()
        for row in rows:
            book.active.append(row)
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    def test_xlsx_values_and_real_excel_date(self):
        from datetime import date
        payload = self.make_xlsx([['外部キー', '商品名', '仕入日', '仕入額'], ['A-1', '仮商品', date(2026, 10, 22), 30000]])
        p = parse_inventory_file(payload, 'test.xlsx')
        self.assertEqual(p['error_count'], 0)
        self.assertEqual(p['rows'][0]['record']['purchase_date'], '2026-10-22')

    def test_xlsx_formula_is_rejected(self):
        payload = self.make_xlsx([['外部キー', '商品名'], ['A', '=1+1']])
        with self.assertRaises(ReconciliationError) as raised:
            parse_inventory_file(payload, 'test.xlsx')
        self.assertEqual(raised.exception.code, 'xlsx_formula')

    def test_xlsx_active_content_and_xml_entities_are_rejected_before_library(self):
        for name, content in [('xl/vbaProject.bin', b'bad'), ('xl/externalLinks/externalLink1.xml', b'<x/>'),
                              ('xl/worksheets/sheet1.xml', b'<!DOCTYPE x [<!ENTITY x "bad">]><x/>'),
                              ('xl/_rels/a.rels', b'<Relationship TargetMode="External"/>'),
                              ('xl/a.xml', '<!DOCTYPE x><x/>'.encode('utf-16'))]:
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, 'w') as archive:
                archive.writestr(name, content)
            with self.subTest(name=name), self.assertRaises(ReconciliationError):
                parse_inventory_file(stream.getvalue(), 'a.xlsx')

    def test_xlsx_multiple_sheets_are_explicitly_rejected(self):
        from openpyxl import Workbook
        book = Workbook()
        book.create_sheet('unreviewed')
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaises(ReconciliationError) as raised:
            parse_inventory_file(stream.getvalue(), 'a.xlsx')
        self.assertEqual(raised.exception.code, 'xlsx_sheet_count')

    def test_zip_bomb_is_rejected_before_workbook_parsing(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/sharedStrings.xml', b'A' * (2 * 1024 * 1024))
        with self.assertRaises(ReconciliationError) as raised:
            parse_inventory_file(stream.getvalue(), 'compressed.xlsx')
        self.assertEqual(raised.exception.code, 'xlsx_compression')

    def test_xlsx_missing_optional_dependency_has_csv_guidance(self):
        data = self.make_xlsx([['外部キー', '商品名'], ['a', '仮商品']])
        with patch.dict('sys.modules', {'openpyxl': None}):
            with self.assertRaises(ReconciliationError) as raised:
                parse_inventory_file(data, 'a.xlsx')
        self.assertEqual(raised.exception.code, 'xlsx_dependency_missing')


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.conn = fixture()
        self.addCleanup(self.conn.close)
        self.registry = Registry(self.conn)
        self.service = ReconciliationService(self.registry)
        self.actor = Actor(1, True)

    def preview(self, value, customer=2, namespace='legacy-export'):
        parsed = parse_inventory_file(csv_bytes(value), 'fixture.csv')
        return self.service.preview(self.actor, customer, namespace, parsed)

    def confirm(self, p):
        return self.service.confirm(self.actor, p['batch_id'], p['preview_hash'], WHEN)

    def test_preview_only_then_link_existing_without_new_item_work_or_fee(self):
        p = self.preview('ID,商品名\n10,既存バッグ\n')
        self.assertEqual(p['summary']['link'], 1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_items').fetchone()[0], 0)
        result = self.confirm(p)
        self.assertEqual(result['linked'], 1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM merchandise').fetchone()[0], 2)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_items').fetchone()[0], 1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_work_records').fetchone()[0], 0)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_item_events').fetchone()[0], 0)
        self.assertEqual(result['charges_created'], 0)
        self.assertEqual(self.service.get_batch(self.actor, p['batch_id'])['status'], 'completed')

    def test_external_new_stock_creates_linked_legacy_and_keeps_unknowns(self):
        p = self.preview('外部キー,商品名,保管場所\nA-1,取込商品,開花\n', namespace='external-sheet')
        result = self.confirm(p)
        self.assertEqual(result['created'], 1)
        legacy = self.conn.execute('SELECT * FROM merchandise WHERE id>11').fetchone()
        self.assertIsNone(legacy['purchase_price'])
        self.assertIsNone(legacy['sale_price'])
        self.assertEqual(legacy['custody_location'], 'unknown')
        profile = self.conn.execute('SELECT * FROM business_item_profiles').fetchone()
        self.assertEqual(json.loads(profile['imported_fields_json'])['custody_location'], 'kaika')
        self.assertIsNone(profile['received_at'])
        self.assertIsNone(profile['registration_completed_at'])
        self.assertIsNone(profile['deadline_at'])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_item_links').fetchone()[0], 1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_work_records').fetchone()[0], 0)

    def test_confirmation_replay_is_noop_and_reimport_different_file_skips(self):
        p = self.preview('外部キー,商品名\nA,商品\n')
        first = self.confirm(p)
        repeated = self.confirm(p)
        self.assertTrue(repeated['repeated'])
        p2 = self.preview('外部キー,商品名\nA,商品\n\n')
        self.assertEqual(p2['summary']['skip'], 1)
        self.assertEqual(self.confirm(p2)['skipped'], 1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_items').fetchone()[0], 1)
        self.assertEqual(len(self.service.history(self.actor, 2)), 2)

    def test_second_batch_prepared_before_first_confirm_is_stale_and_rolls_back(self):
        p = self.preview('外部キー,商品名\nA,商品\n')
        p2 = self.preview('外部キー,商品名\nA,商品\n\n')
        self.confirm(p)
        with self.assertRaises(ReconciliationError) as raised:
            self.confirm(p2)
        self.assertEqual(raised.exception.code, 'source_changed')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_items').fetchone()[0], 1)

    def test_cross_customer_legacy_and_file_owner_mismatch_are_held_without_details(self):
        for source in ('ID,商品名\n11,別顧客の靴\n', '外部キー,商品名,顧客ID\nA,商品,3\n'):
            p = self.preview(source)
            self.assertEqual(p['summary']['hold'], 1)
            with self.assertRaises(ReconciliationError):
                self.confirm(p)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_items').fetchone()[0], 0)

    def test_company_or_unknown_stock_scope_is_held_before_confirmation(self):
        for scope in ('admin', None):
            self.conn.execute('UPDATE merchandise SET scope=? WHERE id=10', (scope,))
            preview = self.preview('ID,商品名\n10,既存バッグ\n')
            self.assertEqual(preview['summary']['hold'], 1)
            with self.assertRaises(ReconciliationError):
                self.confirm(preview)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_item_links').fetchone()[0], 0)

    def test_scope_changed_after_preview_cannot_be_confirmed(self):
        preview = self.preview('ID,商品名\n10,既存バッグ\n')
        self.conn.execute("UPDATE merchandise SET scope='admin' WHERE id=10")
        with self.assertRaises(ReconciliationError) as raised:
            self.confirm(preview)
        self.assertEqual(raised.exception.code, 'source_changed')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM business_item_links').fetchone()[0], 0)

    def test_customer_or_document_only_staff_cannot_read_or_create_previews(self):
        p = self.preview('ID,商品名\n10,既存バッグ\n')
        for actor in (Actor(2), Actor(2, True), Actor(4, True)):
            with self.subTest(actor=actor), self.assertRaises((RegistryError, ReconciliationError)):
                self.service.get_batch(actor, p['batch_id'])

    def test_customer_selection_and_namespace_are_required(self):
        parsed = parse_inventory_file(csv_bytes('外部キー,商品名\nA,商品\n'), 'x.csv')
        for customer, namespace in ((None, 'source'), (1, 'source'), (True, 'source'), (2, '')):
            with self.subTest(customer=customer), self.assertRaises(ReconciliationError):
                self.service.preview(self.actor, customer, namespace, parsed)

    def test_changed_source_cost_or_owner_cannot_confirm_stale_preview(self):
        p = self.preview('ID,商品名\n10,既存バッグ\n')
        self.conn.execute('UPDATE merchandise SET purchase_price=31000 WHERE id=10')
        with self.assertRaises(ReconciliationError) as raised:
            self.confirm(p)
        self.assertEqual(raised.exception.code, 'source_changed')
        renewed = self.preview('ID,商品名\n10,既存バッグ\n')
        self.assertNotEqual(renewed['batch_id'], p['batch_id'])
        self.confirm(renewed)

    def test_changed_file_does_not_overwrite_prior_import(self):
        p = self.preview('外部キー,商品名,仕入額\nA,商品,30000\n')
        self.confirm(p)
        changed = self.preview('外部キー,商品名,仕入額\nA,商品,99999\n')
        self.assertEqual(changed['summary']['hold'], 1)
        self.assertEqual(self.conn.execute('SELECT purchase_price FROM merchandise WHERE id>11').fetchone()[0], 30000)

    def test_tampered_preview_hash_or_database_payload_cannot_execute(self):
        p = self.preview('外部キー,商品名\nA,商品\n')
        with self.assertRaises(ReconciliationError):
            self.service.confirm(self.actor, p['batch_id'], 'wrong', WHEN)
        stored = json.loads(self.conn.execute('SELECT preview_json FROM business_import_batches').fetchone()[0])
        stored['customer_id'] = 3
        self.conn.execute('UPDATE business_import_batches SET preview_json=?', (json.dumps(stored),))
        with self.assertRaises(ReconciliationError):
            self.confirm(p)

    def test_fault_after_legacy_insert_rolls_back_entire_batch(self):
        p = self.preview('外部キー,商品名\nA,商品\n')
        with patch.object(self.registry, 'create_item', side_effect=RuntimeError('synthetic fault')):
            with self.assertRaises(RuntimeError):
                self.confirm(p)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM merchandise').fetchone()[0], 2)
        self.assertEqual(self.service.get_batch(self.actor, p['batch_id'])['status'], 'preview')

    def test_db_unique_source_constraint_is_not_only_a_ui_check(self):
        p = self.preview('外部キー,商品名\nA,商品\n')
        self.confirm(p)
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute('INSERT INTO business_import_sources SELECT * FROM business_import_sources')

    def test_missing_batch_and_naive_time_fail(self):
        with self.assertRaises(ReconciliationError) as raised:
            self.service.get_batch(self.actor, 'missing')
        self.assertEqual(raised.exception.code, 'batch_not_found')
        p = self.preview('外部キー,商品名\nA,商品\n')
        with self.assertRaises(ReconciliationError):
            self.service.confirm(self.actor, p['batch_id'], p['preview_hash'], '2026-10-22')


class MigrationTests(unittest.TestCase):
    def test_postgres_statements_are_portable_and_not_executed(self):
        statements = import_schema_statements()
        self.assertEqual(len(statements), 3)
        self.assertNotIn('AUTOINCREMENT', '\n'.join(statements))
        with self.assertRaises(ReconciliationError):
            migrate_import_schema(None, 'postgres')

    def test_file_database_requires_explicit_temporary_fixture(self):
        with tempfile.TemporaryDirectory() as path:
            conn = sqlite3.connect(str(Path(path) / 'fixture.db'))
            try:
                conn.execute('PRAGMA foreign_keys=ON')
                with self.assertRaises(ReconciliationError):
                    migrate_import_schema(conn)
            finally:
                conn.close()


if __name__ == '__main__':
    unittest.main()
