"""Bounded inventory-file previews and explicit, staff-only reconciliation.

No database, file or network access occurs on import.  Historical stock import
never calls legacy import handlers, creates work logs or calculates charges.
CSV is always supported; XLSX needs the optional, existing openpyxl runtime.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import tempfile
import uuid
import zipfile


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 64
MAX_CELL_CHARS = 4000
PARSER_VERSION = 'inventory-import-v1'

FIELD_LABELS = {
    'legacy_id': '既存システムID', 'external_key': '外部キー（管理番号）',
    'external_label': '既存ラベル',
    'title': '商品名', 'owner_user_id': '顧客ID', 'custody_location': '保管場所',
    'quantity': '数量', 'purchase_date': '仕入日', 'store_name': '店舗名',
    'purchase_price': '仕入額', 'payment_method': '支払方法',
    'listing_price': '出品価格', 'is_listed': '出品済み', 'listing_date': '出品日',
    'sale_date': '売却日', 'sale_price': '売上金', 'shipping_cost': '送料',
    'sales_destination': '販売先', 'commission': '手数料', 'is_shipped': '発送済み',
    'received_at': '入庫日', 'registration_completed_at': '登録完了日', 'shipped_at': '出庫日',
}
_ALIASES = {
    'legacy_id': ('ID', '既存ID', '既存システムID', 'merchandise_id'),
    'external_key': ('外部キー', '管理番号', 'レコード番号', 'external_id', 'source_id'),
    'external_label': ('ラベル', '現物ラベル'),
    'title': ('商品名', '品名', 'product_name'),
    'owner_user_id': ('顧客ID', '所有者ID', 'user_id'),
    'custody_location': ('保管場所', '保管先'), 'quantity': ('数量', '点数'),
    'received_at': ('入庫日', '受領日'), 'registration_completed_at': ('登録完了日', '撮影登録完了日'),
    'shipped_at': ('出庫日', '発送日'),
}
for _field, _label in FIELD_LABELS.items():
    _ALIASES[_field] = tuple(set(_ALIASES.get(_field, ()) + (_label, _field)))
_MONEY_FIELDS = {'purchase_price', 'listing_price', 'sale_price', 'shipping_cost', 'commission'}
_DATE_FIELDS = {'purchase_date', 'listing_date', 'sale_date', 'received_at', 'registration_completed_at', 'shipped_at'}
_BOOL_FIELDS = {'is_listed', 'is_shipped'}
_KNOWN_IGNORED = {'登録者', '利益', '利益率', 'owner_name', 'profit', 'profit_rate'}


class ReconciliationError(ValueError):
    def __init__(self, code, message):
        self.code = code
        self.status = 403 if code == 'staff_required' else 404 if code == 'batch_not_found' else 409 if code in (
            'preview_changed', 'source_changed', 'rows_on_hold') else 400
        super().__init__(message)


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _digest(value):
    return hashlib.sha256(_json(value).encode('utf-8')).hexdigest()


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        if value.time().isoformat() != '00:00:00':
            return value.isoformat()
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def suggest_mapping(headers):
    """Return canonical field -> exact header. Ambiguous aliases stay unmapped."""
    result = {}
    for field, aliases in _ALIASES.items():
        matches = [h for h in headers if h in aliases]
        if len(matches) == 1:
            result[field] = matches[0]
    return result


def _read_csv(data):
    try:
        value = data.decode('utf-8-sig')
        encoding = 'utf-8-sig'
    except UnicodeDecodeError:
        try:
            value = data.decode('cp932')
            encoding = 'cp932'
        except UnicodeDecodeError as exc:
            raise ReconciliationError('encoding', 'UTF-8または日本語Windows形式のCSVを使用してください。') from exc
    if '\x00' in value:
        raise ReconciliationError('binary_content', 'CSVにバイナリ文字が含まれています。')
    try:
        return list(csv.reader(io.StringIO(value), strict=True)), encoding
    except csv.Error as exc:
        raise ReconciliationError('invalid_csv', 'CSVの引用符または列構造を確認してください。') from exc


def _read_xlsx(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 512 or len({i.filename for i in infos}) != len(infos):
                raise ReconciliationError('xlsx_container', 'Excel内部構造が複雑すぎるか重複しています。')
            total = 0
            for entry in infos:
                name = entry.filename.replace('\\', '/')
                lower = name.lower()
                if name.startswith('/') or '..' in name.split('/') or entry.flag_bits & 1:
                    raise ReconciliationError('xlsx_container', '暗号化または不正なExcel内部参照は取り込めません。')
                if any(part in lower for part in ('vbaproject', 'externallinks/', 'embeddings/', 'connections.xml', 'querytables/')):
                    raise ReconciliationError('xlsx_active_content', 'マクロ・外部参照・埋込データを除いたExcelを使用してください。')
                total += entry.file_size
                if total > MAX_UNCOMPRESSED_BYTES or entry.file_size > MAX_UNCOMPRESSED_BYTES:
                    raise ReconciliationError('xlsx_size', '展開後のExcelサイズが上限を超えています。')
                if entry.file_size > 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > 200:
                    raise ReconciliationError('xlsx_compression', '圧縮率が高すぎるExcelは取り込めません。')
                if lower.endswith(('.xml', '.rels')):
                    raw = archive.read(entry)
                    if b'\x00' in raw or b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                        raise ReconciliationError('xlsx_xml', '外部定義を含むExcelは取り込めません。')
                    if lower.endswith('.rels') and re.search(br'TargetMode\s*=\s*[\"\x27]External', raw, re.I):
                        raise ReconciliationError('xlsx_external_reference', '外部リンクを除いたExcelを使用してください。')
    except (zipfile.BadZipFile, RuntimeError) as exc:
        raise ReconciliationError('invalid_xlsx', 'Excelファイルを読み込めません。CSVでも確認できます。') from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ReconciliationError('xlsx_dependency_missing', 'この環境はExcel読込の準備前です。CSV形式で保存して読み込んでください。') from exc
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise ReconciliationError('xlsx_sheet_count', '確認対象だけを1シートにしたExcelを使用してください。')
            sheet = book.worksheets[0]
            if sheet.sheet_state != 'visible' or sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
                raise ReconciliationError('xlsx_dimensions', 'シートの表示状態、行数、列数を確認してください。')
            rows = []
            for row in sheet.iter_rows():
                if any(getattr(cell, 'data_type', None) == 'f' for cell in row):
                    raise ReconciliationError('xlsx_formula', '数式は値に変換してから読み込んでください。')
                rows.append([_cell(cell.value) for cell in row])
            return rows, 'xlsx'
        finally:
            book.close()
    except ReconciliationError:
        raise
    except Exception as exc:
        raise ReconciliationError('invalid_xlsx', 'Excelファイルを読み込めません。形式を確認してください。') from exc


def _validate_value(field, text):
    if not text:
        return None
    if field in _MONEY_FIELDS:
        normal = text.replace(',', '').replace('¥', '').replace('￥', '')
        if not re.fullmatch(r'[0-9]{1,12}', normal):
            raise ValueError('0以上の整数円で入力してください。小数や計算式は使えません。')
        return int(normal)
    if field in ('legacy_id', 'owner_user_id', 'quantity'):
        if not re.fullmatch(r'[0-9]{1,12}', text) or int(text) < 1:
            raise ValueError('1以上の整数で入力してください。')
        if field == 'quantity' and int(text) != 1:
            raise ValueError('1行を1つの現物に分け、異なる管理番号を付けてください。')
        return int(text)
    if field in _DATE_FIELDS:
        normal = text.replace('/', '-')
        try:
            return date.fromisoformat(normal).isoformat()
        except ValueError as exc:
            raise ValueError('日付はYYYY-MM-DDで入力してください。日時や不明値は推測しません。') from exc
    if field in _BOOL_FIELDS:
        if text.lower() in ('済', 'はい', 'true', '1'):
            return True
        if text.lower() in ('未', 'いいえ', 'false', '0'):
            return False
        raise ValueError('済/未、はい/いいえ、1/0のいずれかを指定してください。')
    if field == 'custody_location':
        lookup = {'self': 'self', '自宅': 'self', '自己保管': 'self', 'kaika': 'kaika', '開花': 'kaika',
                  'transit': 'transit', '配送中': 'transit', 'unknown': 'unknown', '不明': 'unknown'}
        if text not in lookup:
            raise ValueError('自宅・開花・配送中・不明のいずれかを指定してください。')
        return lookup[text]
    if field in ('external_key', 'title') and len(text) > 200:
        raise ValueError('200文字以内で入力してください。')
    return text


def parse_inventory_file(data, filename, mapping=None):
    """Pure preview. Mapping is canonical field -> input header, never code."""
    if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
        raise ReconciliationError('file_size', '空ではない5MB以下のファイルを使用してください。')
    extension = str(filename).rsplit('.', 1)[-1].lower()
    if extension == 'csv':
        matrix, encoding = _read_csv(data)
    elif extension == 'xlsx':
        matrix, encoding = _read_xlsx(data)
    else:
        raise ReconciliationError('file_type', 'CSVまたはマクロなしのXLSXを使用してください。')
    if not matrix or len(matrix) > MAX_ROWS + 1:
        raise ReconciliationError('row_limit', '見出しと5,000行以下の商品データを使用してください。')
    headers = [_cell(x) for x in matrix[0]]
    if not headers or len(headers) > MAX_COLUMNS or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ReconciliationError('headers', '見出しは空欄や重複なし、64列以下にしてください。')
    if any(len(h) > 100 for h in headers):
        raise ReconciliationError('headers', '列見出しは100文字以内にしてください。')
    suggested = suggest_mapping(headers)
    if mapping is not None and (not isinstance(mapping, dict) or any(
            not isinstance(key, str) or not isinstance(value, str) for key, value in mapping.items())):
        raise ReconciliationError('mapping', '列対応は項目名と列見出しの組合せで指定してください。')
    selected = dict(suggested if mapping is None else mapping)
    if any(k not in FIELD_LABELS or v not in headers for k, v in selected.items()) or len(set(selected.values())) != len(selected):
        raise ReconciliationError('mapping', '同じ列の重複指定や不明な列指定を確認してください。')
    warnings = []
    ignored = [h for h in headers if h not in selected.values()]
    if ignored:
        warnings.append('未使用列: ' + '、'.join(ignored) + '。利益・料金・登録者名から所有者や請求額を推測しません。')
    rows = []
    for number, source in enumerate(matrix[1:], start=2):
        values = [_cell(x) for x in source]
        if not any(values):
            continue
        errors = []
        if len(values) != len(headers):
            errors.append('列数が見出しと一致しません。')
        if any(len(v) > MAX_CELL_CHARS for v in values):
            errors.append('1セルは4,000文字以内にしてください。')
        if any(v.lstrip().startswith(('=', '@', '+', '-')) for v in values):
            errors.append('計算式・外部関数の文字列は値に変換してください。')
        raw = dict(zip(headers, values))
        clean = {key: None for key in FIELD_LABELS}
        for key, header in selected.items():
            try:
                clean[key] = _validate_value(key, raw.get(header, ''))
            except ValueError as exc:
                errors.append(FIELD_LABELS[key] + ': ' + str(exc))
        if not clean['title']:
            errors.append('商品名を指定してください。')
        if not clean['legacy_id'] and not clean['external_key']:
            errors.append('既存システムIDまたは外部キーが必要です。行番号や商品名だけでは同一商品を確定しません。')
        clean['custody_location'] = clean['custody_location'] or 'unknown'
        rows.append({'row_number': number, 'record': clean, 'errors': errors,
                     'warnings': ['未入力値は不明として保存し、0円・作業日・保管期限を補いません。']})
    if not rows:
        raise ReconciliationError('empty_records', '商品データがありません。')
    groups = {}
    for row in rows:
        record = row['record']
        for field in ('legacy_id', 'external_key'):
            if record[field] is not None:
                groups.setdefault(field + ':' + str(record[field]), []).append(row)
    for duplicates in groups.values():
        if len(duplicates) > 1:
            for row in duplicates:
                row['errors'].append('ファイル内で同じ識別子が重複しています。自動統合しません。')
    return {'parser_version': PARSER_VERSION, 'file_sha256': hashlib.sha256(data).hexdigest(),
            'encoding': encoding, 'headers': headers, 'mapping': selected, 'suggested_mapping': suggested,
            'unused_headers': ignored, 'warnings': warnings, 'rows': rows,
            'error_count': sum(bool(r['errors']) for r in rows), 'row_count': len(rows),
            'creates_work_records': False, 'creates_charges': False}


def import_schema_statements():
    """Additive portable SQL. Return text only; no production application."""
    return (
        '''CREATE TABLE IF NOT EXISTS business_import_batches (
            batch_id TEXT PRIMARY KEY, customer_id INTEGER NOT NULL REFERENCES users(id),
            actor_id INTEGER NOT NULL REFERENCES users(id), source_namespace TEXT NOT NULL,
            file_sha256 TEXT NOT NULL, mapping_sha256 TEXT NOT NULL,
            preview_hash TEXT NOT NULL, preview_json TEXT NOT NULL,
            status TEXT NOT NULL CHECK(status IN ('preview','applying','completed','blocked')),
            created_at TEXT NOT NULL, completed_at TEXT, result_json TEXT,
            UNIQUE(customer_id, source_namespace, file_sha256, mapping_sha256, preview_hash)
        )''',
        '''CREATE TABLE IF NOT EXISTS business_import_sources (
            customer_id INTEGER NOT NULL REFERENCES users(id), source_namespace TEXT NOT NULL,
            external_key TEXT NOT NULL, physical_item_id TEXT NOT NULL REFERENCES business_items(physical_item_id),
            record_sha256 TEXT NOT NULL, batch_id TEXT NOT NULL REFERENCES business_import_batches(batch_id),
            row_number INTEGER NOT NULL CHECK(row_number>=2),
            PRIMARY KEY(customer_id, source_namespace, external_key)
        )''',
        '''CREATE TABLE IF NOT EXISTS business_import_rows (
            batch_id TEXT NOT NULL REFERENCES business_import_batches(batch_id),
            row_number INTEGER NOT NULL CHECK(row_number>=2), external_key TEXT NOT NULL,
            action TEXT NOT NULL CHECK(action IN ('create','link','skip','hold')),
            physical_item_id TEXT REFERENCES business_items(physical_item_id),
            record_sha256 TEXT NOT NULL, errors_json TEXT NOT NULL,
            PRIMARY KEY(batch_id,row_number)
        )''',
    )


def migrate_import_schema(connection, dialect='sqlite', allow_temporary=False):
    """Explicit fixture-only application; PostgreSQL SQL is reviewable only."""
    if dialect != 'sqlite':
        raise ReconciliationError('migration_review_required', 'PostgreSQLへの適用は移行レビュー後に行ってください。')
    paths = [row[2] for row in connection.execute('PRAGMA database_list').fetchall() if row[2]]
    for value in paths:
        path = Path(value).resolve()
        if not allow_temporary or not path.is_relative_to(Path(tempfile.gettempdir()).resolve()):
            raise ReconciliationError('fixture_database_required', '取込スキーマの適用先は明示した使い捨て検証DBだけです。')
    if connection.execute('PRAGMA foreign_keys').fetchone()[0] != 1:
        raise ReconciliationError('foreign_keys_required', '外部キー制約を有効にしてください。')
    connection.execute('SAVEPOINT import_schema')
    try:
        for statement in import_schema_statements():
            connection.execute(statement)
        connection.execute('RELEASE SAVEPOINT import_schema')
    except Exception:
        connection.execute('ROLLBACK TO SAVEPOINT import_schema')
        connection.execute('RELEASE SAVEPOINT import_schema')
        raise


class ReconciliationService:
    """Uses Registry transactions. Caller commits its own isolated connection."""

    def __init__(self, registry):
        self.registry = registry

    def _authorize(self, actor, customer_id=None):
        self.registry.actor(actor, staff_required=True)
        users = self.registry.query('SELECT id,role FROM users WHERE id=?', (actor.user_id,))
        if not getattr(actor, 'staff', False) or not users or users[0]['role'] not in ('owner', 'admin'):
            raise ReconciliationError('staff_required', '在庫照合は権限のあるスタッフが行ってください。')
        if customer_id is not None:
            if isinstance(customer_id, bool) or not isinstance(customer_id, int) or customer_id <= 0:
                raise ReconciliationError('customer_required', '対象の顧客を選択してください。')
            customers = self.registry.query('SELECT id,role FROM users WHERE id=?', (customer_id,))
            if not customers or customers[0]['role'] != 'user':
                raise ReconciliationError('customer_required', '確認済みの対象顧客を選択してください。')

    def _source_key(self, record):
        if record.get('legacy_id'):
            return 'legacy:' + str(record['legacy_id'])
        return 'external:' + str(record.get('external_key') or '')

    def _plan_row(self, customer_id, namespace, row):
        record = row['record']
        errors = list(row['errors'])
        key = self._source_key(record)
        result = {'row_number': row['row_number'], 'external_key': key, 'record': record,
                  'record_sha256': _digest(record), 'action': 'create', 'physical_item_id': None,
                  'errors': errors, 'warnings': list(row.get('warnings', [])), 'legacy_fingerprint': None}
        if record.get('owner_user_id') not in (None, customer_id):
            errors.append('ファイルの顧客IDが選択した顧客と一致しません。')
        existing = self.registry.query('''SELECT physical_item_id,record_sha256 FROM business_import_sources
            WHERE customer_id=? AND source_namespace=? AND external_key=?''', (customer_id, namespace, key))
        if existing:
            result['physical_item_id'] = existing[0]['physical_item_id']
            if existing[0]['record_sha256'] == result['record_sha256']:
                result['action'] = 'skip'
            else:
                errors.append('以前の取込と内容が異なります。既存商品を上書きせず訂正確認をしてください。')
        elif record.get('legacy_id'):
            legacy = self.registry.query('SELECT * FROM merchandise WHERE id=?', (record['legacy_id'],))
            if not legacy or legacy[0]['user_id'] != customer_id:
                errors.append('対象顧客の既存IDとして照合できません。他顧客の詳細は表示しません。')
            elif 'scope' in legacy[0] and legacy[0]['scope'] != 'user':
                errors.append('顧客の自己・受託在庫として確認できません。開花の自社商品や区分不明の商品は別途照合してください。')
            elif legacy[0]['product_name'] != record['title']:
                errors.append('既存IDの商品名が一致しません。確認してから再度読み込んでください。')
            else:
                result['legacy_fingerprint'] = _digest({k: _cell(v) for k, v in legacy[0].items()})
                links = self.registry.query('SELECT physical_item_id FROM business_item_links WHERE merchandise_id=?', (record['legacy_id'],))
                if links:
                    result['physical_item_id'] = links[0]['physical_item_id']
                result['action'] = 'link'
        if errors:
            result['action'] = 'hold'
        return result

    def preview(self, actor, customer_id, source_namespace, parsed):
        if type(customer_id) is not int or customer_id <= 0:
            raise ReconciliationError('customer_required', '対象の顧客を選択してください。')
        self._authorize(actor, customer_id)
        namespace = str(source_namespace or '').strip()
        if not re.fullmatch(r'[A-Za-z0-9_.:-]{1,80}', namespace):
            raise ReconciliationError('source_namespace', '取込元名は英数字・ハイフン等で80文字以内にしてください。例: legacy-export')
        if not isinstance(parsed, dict) or parsed.get('parser_version') != PARSER_VERSION:
            raise ReconciliationError('preview_required', 'ファイルを読み直してプレビューしてください。')
        # HTTP must pass the server parser result; do not accept client JSON as a
        # parsed file.  Validate again to protect internal callers as well.
        if not isinstance(parsed.get('rows'), list) or not 0 < len(parsed['rows']) <= MAX_ROWS:
            raise ReconciliationError('preview_required', '取込行数を確認してください。')
        mapping_hash = _digest(parsed.get('mapping'))
        rows = [self._plan_row(customer_id, namespace, row) for row in parsed['rows']]
        summary = {action: sum(row['action'] == action for row in rows) for action in ('create', 'link', 'skip', 'hold')}
        plan = {'customer_id': customer_id, 'source_namespace': namespace, 'file_sha256': parsed['file_sha256'],
                'mapping': parsed['mapping'], 'rows': rows, 'summary': summary,
                'warnings': parsed.get('warnings', []), 'creates_work_records': False, 'creates_charges': False}
        preview_hash = _digest(plan)
        prior = self.registry.query('''SELECT * FROM business_import_batches WHERE customer_id=? AND source_namespace=?
            AND file_sha256=? AND mapping_sha256=? AND preview_hash=?''',
            (customer_id, namespace, parsed['file_sha256'], mapping_hash, preview_hash))
        if prior:
            return self._batch_result(prior[0])
        batch_id = str(uuid.uuid4())
        stamp = datetime.now(timezone.utc).isoformat()
        with self.registry.atomic():
            self.registry.execute('''INSERT INTO business_import_batches
                (batch_id,customer_id,actor_id,source_namespace,file_sha256,mapping_sha256,preview_hash,preview_json,status,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?)''',
                (batch_id, customer_id, actor.user_id, namespace, parsed['file_sha256'], mapping_hash,
                 preview_hash, _json(plan), 'preview', stamp))
            for row in rows:
                self.registry.execute('''INSERT INTO business_import_rows
                    (batch_id,row_number,external_key,action,physical_item_id,record_sha256,errors_json) VALUES(?,?,?,?,?,?,?)''',
                    (batch_id, row['row_number'], row['external_key'], row['action'], row['physical_item_id'],
                     row['record_sha256'], _json(row['errors'])))
        return dict(plan, batch_id=batch_id, preview_hash=preview_hash, status='preview', created_at=stamp)

    def _batch_result(self, batch):
        result = json.loads(batch['preview_json'])
        result.update(batch_id=batch['batch_id'], preview_hash=batch['preview_hash'], status=batch['status'], created_at=batch['created_at'])
        if batch.get('result_json'):
            result['result'] = json.loads(batch['result_json'])
        return result

    def get_preview(self, actor, batch_id):
        self._authorize(actor)
        found = self.registry.query('SELECT * FROM business_import_batches WHERE batch_id=?', (batch_id,))
        if not found:
            raise ReconciliationError('batch_not_found', '取込プレビューが見つかりません。')
        self._authorize(actor, found[0]['customer_id'])
        return self._batch_result(found[0])

    get_batch = get_preview

    def confirm(self, actor, batch_id, preview_hash, occurred_at):
        self._authorize(actor)
        if not isinstance(occurred_at, str):
            raise ReconciliationError('occurred_at', '時刻はタイムゾーン付きで指定してください。')
        try:
            parsed_time = datetime.fromisoformat(occurred_at.replace('Z', '+00:00'))
            if parsed_time.tzinfo is None:
                raise ValueError()
        except ValueError as exc:
            raise ReconciliationError('occurred_at', '時刻はタイムゾーン付きで指定してください。') from exc
        with self.registry.atomic():
            suffix = ' FOR UPDATE' if getattr(self.registry, 'dialect', 'sqlite') == 'postgres' else ''
            found = self.registry.query('SELECT * FROM business_import_batches WHERE batch_id=?' + suffix, (batch_id,))
            if not found:
                raise ReconciliationError('batch_not_found', '取込プレビューが見つかりません。')
            batch = found[0]
            self._authorize(actor, batch['customer_id'])
            if batch['preview_hash'] != preview_hash or _digest(json.loads(batch['preview_json'])) != preview_hash:
                raise ReconciliationError('preview_changed', 'プレビューが変更されています。読み直してください。')
            if batch['status'] == 'completed':
                return dict(json.loads(batch['result_json']), repeated=True)
            plan = json.loads(batch['preview_json'])
            if batch['status'] != 'preview' or plan['summary']['hold']:
                raise ReconciliationError('rows_on_hold', '保留行を修正し、全行が照合できてから確定してください。')
            for row in plan['rows']:
                current = self._plan_row(batch['customer_id'], batch['source_namespace'],
                                         dict(row, errors=[], warnings=row['warnings']))
                compared = ('action', 'physical_item_id', 'record_sha256', 'legacy_fingerprint')
                if any(current[key] != row[key] for key in compared) or current['errors']:
                    raise ReconciliationError('source_changed', 'プレビュー後に照合先が変わりました。最新状態で再確認してください。')
            self.registry.execute("UPDATE business_import_batches SET status='applying' WHERE batch_id=? AND status='preview'", (batch_id,))
            counts = {'created': 0, 'linked': 0, 'skipped': 0, 'work_records_created': 0, 'charges_created': 0}
            for row in plan['rows']:
                if row['action'] == 'skip':
                    counts['skipped'] += 1
                    continue
                record = row['record']
                source_key = 'import:' + _digest([batch['customer_id'], batch['source_namespace'], row['external_key']])
                evidence = 'import-batch:' + batch_id + ':row:' + str(row['row_number'])
                if row['action'] == 'link':
                    item = self.registry.attach_legacy(actor, record['legacy_id'], physical_item_id=row['physical_item_id'],
                        evidence_reference=evidence, occurred_at=occurred_at, source_key=source_key)
                    counts['linked'] += 1
                else:
                    imported = {k: v for k, v in record.items() if k not in ('title', 'owner_user_id', 'quantity', 'external_label')}
                    legacy_id = self._insert_legacy(batch['customer_id'], record)
                    item = self.registry.create_item(actor, title=record['title'], owner_user_id=batch['customer_id'],
                        custodian='unknown', occurred_at=occurred_at, evidence_reference=evidence,
                        source_key=source_key, origin='legacy_import', imported_fields=imported,
                        external_label=record.get('external_label'))
                    self.registry.attach_legacy(actor, legacy_id, physical_item_id=item['physical_item_id'],
                        occurred_at=occurred_at, evidence_reference=evidence, source_key=source_key + ':legacy')
                    counts['created'] += 1
                item_id = item['physical_item_id']
                self.registry.execute('''INSERT INTO business_import_sources
                    (customer_id,source_namespace,external_key,physical_item_id,record_sha256,batch_id,row_number) VALUES(?,?,?,?,?,?,?)''',
                    (batch['customer_id'], batch['source_namespace'], row['external_key'], item_id,
                     row['record_sha256'], batch_id, row['row_number']))
                self.registry.execute('UPDATE business_import_rows SET physical_item_id=? WHERE batch_id=? AND row_number=?',
                                      (item_id, batch_id, row['row_number']))
            counts.update(batch_id=batch_id, completed_at=occurred_at, repeated=False)
            self.registry.execute("UPDATE business_import_batches SET status='completed',completed_at=?,result_json=? WHERE batch_id=?",
                                  (occurred_at, _json(counts), batch_id))
            return counts

    def _insert_legacy(self, customer_id, record):
        """Same transaction, allowlisted raw data only; no legacy work hooks."""
        if self.registry.dialect == 'postgres':
            columns = {r['column_name'] for r in self.registry.query(
                "SELECT column_name FROM information_schema.columns WHERE table_schema=current_schema() AND table_name='merchandise'")}
        else:
            columns = {r['name'] for r in self.registry.query('PRAGMA table_info(merchandise)')}
        fields = {'user_id': customer_id, 'product_name': record['title']}
        for field in ('purchase_date', 'store_name', 'purchase_price', 'payment_method', 'listing_price', 'is_listed',
                      'listing_date', 'sale_date', 'sale_price', 'shipping_cost', 'sales_destination', 'commission', 'is_shipped'):
            if field in columns:
                fields[field] = record.get(field)
        # Do not treat historical source locations as receipt evidence. The
        # original value remains in imported_fields for staff reconciliation.
        if 'custody_location' in columns:
            fields['custody_location'] = 'unknown'
        if 'scope' in columns:
            fields['scope'] = 'user'
        names = list(fields)
        sql = 'INSERT INTO merchandise (' + ','.join(names) + ') VALUES (' + ','.join('?' for _ in names) + ')'
        if self.registry.dialect == 'postgres':
            return self.registry.query(sql + ' RETURNING id', tuple(fields[n] for n in names))[0]['id']
        cursor = self.registry.execute(sql, tuple(fields[n] for n in names))
        return cursor.lastrowid

    def history(self, actor, customer_id):
        if type(customer_id) is not int or customer_id <= 0:
            raise ReconciliationError('customer_required', '対象の顧客を選択してください。')
        self._authorize(actor, customer_id)
        batches = self.registry.query('''SELECT batch_id,source_namespace,status,created_at,completed_at,result_json
            FROM business_import_batches WHERE customer_id=? ORDER BY created_at DESC''', (customer_id,))
        for batch in batches:
            batch['result'] = json.loads(batch.pop('result_json')) if batch.get('result_json') else None
        return batches

